from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"D:\VS_project\MyWeb\output\马浩_职业生涯规划比赛简历.docx")


def set_font(run, size=9.3, bold=False, color="222222"):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return run


def add_text(p, text, *, size=9.3, bold=False, color="222222"):
    return set_font(p.add_run(text), size=size, bold=bold, color=color)


def add_link(p, text, url):
    part = p.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "222222")
    rpr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.append(size)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    p._p.append(hyperlink)


doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7)
sec.page_width = Cm(21)
sec.top_margin = Cm(1.55)
sec.bottom_margin = Cm(1.45)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)
sec.header_distance = Cm(0.7)
sec.footer_distance = Cm(0.7)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(9.3)
normal.font.color.rgb = RGBColor(34, 34, 34)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(0)

title_style = styles["Title"]
title_style.font.name = "Microsoft YaHei"
title_style.font.size = Pt(20)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(2)
title_style.paragraph_format.keep_with_next = True
title_ppr = title_style.element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)


def section(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_text(p, title, size=10.5, bold=True, color="000000")
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "C9CED3")
    borders.append(bottom)
    ppr.append(borders)
    return p


def item_head(title, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    add_text(p, title, size=9.5, bold=True, color="000000")
    if meta:
        add_text(p, "  |  " + meta, size=8.7, color="555555")
    return p


def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(2.5)
    add_text(p, "•  " + text, size=9.2)
    return p


p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(p, "马浩", size=20, bold=True, color="000000")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
add_text(p, "职业目标  光电感知与智能算法工程师", size=10.2, bold=True, color="000000")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
add_text(p, "苏州 · ", size=9)
add_link(p, "2030985559@qq.com", "mailto:2030985559@qq.com")
add_text(p, " · ", size=9)
add_link(p, "github.com/MaHao777", "https://github.com/MaHao777")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
add_text(p, "聚焦事件视觉与高速光电感知，持续把物理建模、机器学习和可运行的软件系统结合起来。", size=9.2)

section("教育背景")
item_head("苏州大学 · 光电科学与工程学院", "本科在读 · 2024 级集萃行本实验班")
bullet("2026 年苏州大学“莙政基金”项目入选；大学英语六级 606 分、四级 623 分。")

section("科研与项目实践")
item_head("事件视觉与波前数据平台", "2026 · 训练数据软件开发")
bullet("使用 Python、PySide6 开发训练数据生成平台，连接湍流波前、哈特曼光斑与事件响应；实现五视图预览、批量参数扫描、HDF5 录制和回放。")
bullet("团队作品《极速视界》获第十四届全国大学生光电设计竞赛（东部区赛）一等奖；本人负责数据软件，并参与视频制作与现场答辩。")

item_head("TraceFormer 事件小目标检测", "2026 · 模型与评测")
bullet("设计并训练事件流检测模型，结合 LightGBM 后处理与轨迹规则，完成推理、消融实验和原始事件预测映射核验。")
bullet("冻结完整系统在 24 个验证序列的本地独立复评得分 0.9541；该指标为验证集结果，不代表官方隐藏测试成绩。")

item_head("干涉条纹闭环锁定", "2026 · 光电实验与控制")
bullet("参与线阵相机接入和标定，开发 MATLAB 控制程序串联相位解调与压电补偿；完成整体标定和循环控制验证。")

section("企业工程实践")
item_head("智慧作业车辆通信开发", "清华大学苏州汽车研究院相关企业实习 · 约四个月")
bullet("开发双向 MQTT 模拟测试工具，支持连接、消息收发与日志记录；编写 ROS 2 消息中转节点，处理网页端与 CAN 端通信并完成实测上线。")
bullet("参与不同机型的程序适配、现场联调与缺陷排查，形成从问题定位到验证修复的工程闭环。")

section("竞赛荣誉与技能")
bullet("第十六届蓝桥杯单片机设计与开发大学组：江苏赛区一等奖、全国总决赛优秀奖（2025）；LuminaMind 获思必驰“智会未来”OPC 创新大赛优秀项目奖（团队，2026）。")
bullet("技术：Python、PyTorch、MATLAB、C++、Qt / PySide6、ROS 2、MQTT、Git、Linux；研究主题：事件相机、波前重构、光电控制。")

doc.core_properties.title = "马浩 职业生涯规划比赛简历"
doc.core_properties.subject = "光电感知与智能算法方向"
doc.core_properties.author = "马浩"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
